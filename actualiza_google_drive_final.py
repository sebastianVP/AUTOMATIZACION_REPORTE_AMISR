# -*- coding: utf-8 -*-

import io
import pandas as pd

from openpyxl import load_workbook

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload


# ============================================================
# CONFIGURACIÓN
# ============================================================

URL_CSV = (
    "https://raw.githubusercontent.com/sebastianVP/"
    "DATASETS_CLASE/refs/heads/main/REPORTE_DIGDT_2026.csv"
)

# Credenciales fuera del repositorio
CREDENTIALS = "/home/soporte/.google/amisr-service-account.json"

# ID del archivo XLSX en Google Drive
FILE_ID = "1NQfdeLyCrD97uTQYo3BttK26JYXuaRzi"

SCOPES = ["https://www.googleapis.com/auth/drive"]


# ============================================================
# CONFIGURACIÓN DE MESES Y COLUMNAS
# ============================================================

MESES_COLUMNAS = {
    "AGOSTO": "J",
    "SEPTIEMBRE": "K",
    "OCTUBRE": "L",
    "NOVIEMBRE": "M",
    "DICIEMBRE": "N"
}


# ============================================================
# LEER REPORTE DESDE GITHUB
# ============================================================

def leer_reporte():

    df = pd.read_csv(URL_CSV)

    df["MES"] = df["MES"].astype(str).str.strip().str.upper()
    df["TIPO"] = df["TIPO"].astype(str).str.strip().str.upper()
    df["HORAS"] = pd.to_numeric(
        df["HORAS"],
        errors="coerce"
    ).fillna(0)

    return df


# ============================================================
# OBTENER HORAS POR MES Y TIPO
# ============================================================

def obtener_horas(df, mes, tipo):

    datos = df[
        (df["MES"] == mes) &
        (df["TIPO"] == tipo)
    ]

    return round(datos["HORAS"].sum(), 2)


# ============================================================
# BUSCAR HOJA POR TEXTO
# ============================================================

def obtener_hoja(wb, texto):

    for nombre in wb.sheetnames:

        if texto.lower() in nombre.lower():

            print(f"   ✓ Hoja encontrada: {nombre}")

            return wb[nombre]

    raise ValueError(
        f"No se encontró una hoja que contenga: {texto}"
    )


# ============================================================
# CONECTAR CON GOOGLE DRIVE
# ============================================================

def conectar_drive():

    credenciales = Credentials.from_service_account_file(
        CREDENTIALS,
        scopes=SCOPES
    )

    return build(
        "drive",
        "v3",
        credentials=credenciales
    )


# ============================================================
# OBTENER INFORMACIÓN DEL ARCHIVO
# ============================================================

def obtener_nombre_archivo(drive):

    archivo = drive.files().get(
        fileId=FILE_ID,
        fields="id,name,mimeType"
    ).execute()

    print(f"📄 Archivo: {archivo['name']}")
    print(f"🆔 ID: {archivo['id']}")
    print(f"📦 Tipo: {archivo['mimeType']}")

    return archivo["name"]


# ============================================================
# DESCARGAR XLSX DESDE GOOGLE DRIVE
# ============================================================

def descargar_excel(drive):

    solicitud = drive.files().get_media(
        fileId=FILE_ID
    )

    archivo = io.BytesIO()

    descargador = MediaIoBaseDownload(
        archivo,
        solicitud
    )

    terminado = False

    while not terminado:

        _, terminado = descargador.next_chunk()

    archivo.seek(0)

    return archivo


# ============================================================
# ACTUALIZAR EXCEL
# ============================================================

def actualizar_excel(archivo, horas):

    wb = load_workbook(archivo)

    print("\n📑 HOJAS ENCONTRADAS EN EL XLSX:")

    for hoja in wb.sheetnames:
        print(f"   - {hoja}")

    # --------------------------------------------------------
    # IDENTIFICAR HOJAS
    # --------------------------------------------------------

    dcg = obtener_hoja(
        wb,
        "Datos geofísicos válidos - DCG"
    )

    digdt = obtener_hoja(
        wb,
        "Instrumentación Ionosférica"
    )

    # --------------------------------------------------------
    # ACTUALIZAR MESES
    # --------------------------------------------------------

    print("\n📝 ACTUALIZANDO CELDAS:")

    for mes, columna in MESES_COLUMNAS.items():

        # ----------------------------------------------------
        # DCG
        # ----------------------------------------------------

        celda_isr = f"{columna}14"
        celda_esf = f"{columna}15"

        dcg[celda_isr] = horas[f"{mes}_ISR"]
        dcg[celda_esf] = horas[f"{mes}_ESF"]

        # ----------------------------------------------------
        # DIGDT
        # ----------------------------------------------------

        celda_total = f"{columna}13"

        digdt[celda_total] = round(
            horas[f"{mes}_ISR"] +
            horas[f"{mes}_ESF"],
            2
        )

        # ----------------------------------------------------
        # MOSTRAR RESULTADOS
        # ----------------------------------------------------

        print(
            f"   ✓ DCG!{celda_isr} → "
            f"{mes}_ISR = "
            f"{horas[f'{mes}_ISR']:.2f}"
        )

        print(
            f"   ✓ DCG!{celda_esf} → "
            f"{mes}_ESF = "
            f"{horas[f'{mes}_ESF']:.2f}"
        )

        print(
            f"   ✓ DIGDT!{celda_total} → "
            f"{mes} TOTAL = "
            f"{digdt[celda_total].value:.2f}"
        )

    # --------------------------------------------------------
    # GUARDAR XLSX EN MEMORIA
    # --------------------------------------------------------

    salida = io.BytesIO()

    wb.save(salida)

    salida.seek(0)

    return salida


# ============================================================
# SUBIR XLSX ACTUALIZADO A GOOGLE DRIVE
# ============================================================

def subir_excel(drive, archivo):

    metadata = {
        "name": obtener_nombre_archivo(drive)
    }

    media = MediaIoBaseUpload(
        archivo,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        resumable=True
    )

    drive.files().update(
        fileId=FILE_ID,
        body=metadata,
        media_body=media
    ).execute()


# ============================================================
# PROCESO PRINCIPAL
# ============================================================

def actualizar():

    print("\n" + "=" * 70)
    print(" ACTUALIZACIÓN REPORTE AMISR-14 - GOOGLE DRIVE")
    print("=" * 70)

    # --------------------------------------------------------
    # 1. LEER CSV DESDE GITHUB
    # --------------------------------------------------------

    print("\n📥 Leyendo reporte desde GitHub...")

    df = leer_reporte()

    # --------------------------------------------------------
    # 2. GENERAR HORAS DINÁMICAMENTE
    # --------------------------------------------------------

    horas = {
        f"{mes}_{tipo}": obtener_horas(df, mes, tipo)
        for mes in MESES_COLUMNAS
        for tipo in ["ISR", "ESF"]
    }

    # --------------------------------------------------------
    # 3. MOSTRAR DATOS OBTENIDOS
    # --------------------------------------------------------

    print("\n📊 DATOS OBTENIDOS:")

    for clave, valor in horas.items():

        print(
            f"   {clave:<20}: "
            f"{valor:.2f} horas"
        )

    # --------------------------------------------------------
    # 4. CONECTAR CON GOOGLE DRIVE
    # --------------------------------------------------------

    print("\n☁️ Conectando con Google Drive...")

    drive = conectar_drive()

    nombre = obtener_nombre_archivo(drive)

    # --------------------------------------------------------
    # 5. DESCARGAR XLSX
    # --------------------------------------------------------

    print("\n📥 Descargando archivo XLSX...")

    archivo = descargar_excel(drive)

    print("   ✓ Archivo descargado correctamente")

    # --------------------------------------------------------
    # 6. MODIFICAR XLSX
    # --------------------------------------------------------

    print("\n📝 Modificando XLSX...")

    archivo_actualizado = actualizar_excel(
        archivo,
        horas
    )

    # --------------------------------------------------------
    # 7. SUBIR XLSX
    # --------------------------------------------------------

    print("\n📤 Subiendo XLSX actualizado a Google Drive...")

    subir_excel(
        drive,
        archivo_actualizado
    )

    # --------------------------------------------------------
    # 8. FINALIZAR
    # --------------------------------------------------------

    print("\n" + "=" * 70)
    print(" ✅ ARCHIVO ACTUALIZADO CORRECTAMENTE")
    print("=" * 70)

    print(f"\n📄 {nombre}")

    print("\n📌 MESES ACTUALIZADOS:")

    for mes, columna in MESES_COLUMNAS.items():

        print(
            f"   {mes:<12} → columna {columna}"
        )

    print(
        "\nLas demás celdas del archivo "
        "no fueron modificadas."
    )


# ============================================================
# EJECUCIÓN
# ============================================================

if __name__ == "__main__":

    actualizar()