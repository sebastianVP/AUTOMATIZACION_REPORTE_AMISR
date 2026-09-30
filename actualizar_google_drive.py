# -*- coding: utf-8 -*-

import io
import pandas as pd
from openpyxl import load_workbook
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

URL_CSV = "https://raw.githubusercontent.com/sebastianVP/DATASETS_CLASE/refs/heads/main/REPORTE_DIGDT_2026.csv"

CREDENTIALS = "/home/soporte/.google/amisr-service-account.json"

# ID DEL ARCHIVO XLSX EN GOOGLE DRIVE
FILE_ID = "1NQfdeLyCrD97uTQYo3BttK26JYXuaRzi"

SCOPES = ["https://www.googleapis.com/auth/drive"]

HOJA_DCG = "Datos geofísicos válidos - DCG"
HOJA_DIGDT = "Instrumentación Ionosférica (Ge"


def leer_reporte():
    df = pd.read_csv(URL_CSV)

    df["MES"] = df["MES"].astype(str).str.strip().str.upper()
    df["TIPO"] = df["TIPO"].astype(str).str.strip().str.upper()
    df["HORAS"] = pd.to_numeric(df["HORAS"], errors="coerce").fillna(0)

    return df


def obtener_horas(df, mes, tipo):
    datos = df[(df["MES"] == mes) & (df["TIPO"] == tipo)]
    return round(datos["HORAS"].sum(), 2)

def obtener_hoja(wb, texto):
    for nombre in wb.sheetnames:
        if texto.lower() in nombre.lower():
            print(f"   ✓ Hoja encontrada: {nombre}")
            return wb[nombre]
    raise ValueError(f"No se encontró una hoja que contenga: {texto}")

def conectar_drive():
    credenciales = Credentials.from_service_account_file(
        CREDENTIALS,
        scopes=SCOPES
    )

    return build("drive", "v3", credentials=credenciales)


def descargar_excel(drive):
    solicitud = drive.files().get_media(fileId=FILE_ID)

    archivo = io.BytesIO()
    descargador = MediaIoBaseDownload(archivo, solicitud)

    terminado = False

    while not terminado:
        _, terminado = descargador.next_chunk()

    archivo.seek(0)

    return archivo


def obtener_nombre_archivo(drive):
    archivo = drive.files().get(
        fileId=FILE_ID,
        fields="id,name,mimeType"
    ).execute()

    print(f"📄 Archivo: {archivo['name']}")
    print(f"🆔 ID: {archivo['id']}")
    print(f"📦 Tipo: {archivo['mimeType']}")

    return archivo["name"]


def actualizar_excel(archivo, horas):
    wb = load_workbook(archivo)

    print("\n📑 HOJAS ENCONTRADAS EN EL XLSX:")
    for hoja in wb.sheetnames:
        print(f"   - {hoja}")

    dcg = obtener_hoja(wb, "Datos geofísicos válidos - DCG")
    digdt = obtener_hoja(wb, "Instrumentación Ionosférica")

    # -------------------------------------------------
    # DATOS GEOFÍSICOS VÁLIDOS - DCG
    # -------------------------------------------------

    dcg["J14"] = horas["AGOSTO_ISR"]
    dcg["K14"] = horas["SEPTIEMBRE_ISR"]
    dcg["L14"] = horas["OCTUBRE_ISR"]

    dcg["J15"] = horas["AGOSTO_ESF"]
    dcg["K15"] = horas["SEPTIEMBRE_ESF"]

    # -------------------------------------------------
    # INSTRUMENTACIÓN IONOSFÉRICA - DIGDT
    # -------------------------------------------------

    digdt["J13"] = round(
        horas["AGOSTO_ISR"] +
        horas["AGOSTO_ESF"],
        2
    )

    digdt["K13"] = round(
        horas["SEPTIEMBRE_ISR"] +
        horas["SEPTIEMBRE_ESF"],
        2
    )

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)

    return salida

def subir_excel(drive, archivo):
    metadata = {
        "name": obtener_nombre_archivo(drive)
    }

    media = MediaIoBaseUpload(
        archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        resumable=True
    )

    drive.files().update(
        fileId=FILE_ID,
        body=metadata,
        media_body=media
    ).execute()


def actualizar():
    print("\n" + "=" * 70)
    print(" ACTUALIZACIÓN REPORTE AMISR-14 - GOOGLE DRIVE")
    print("=" * 70)

    # -------------------------------------------------
    # 1. LEER CSV DESDE GITHUB
    # -------------------------------------------------

    print("\n📥 Leyendo reporte desde GitHub...")

    df = leer_reporte()

    horas = {
        "AGOSTO_ISR": obtener_horas(df, "AGOSTO", "ISR"),
        "AGOSTO_ESF": obtener_horas(df, "AGOSTO", "ESF"),
        "SEPTIEMBRE_ISR": obtener_horas(df, "SEPTIEMBRE", "ISR"),
        "SEPTIEMBRE_ESF": obtener_horas(df, "SEPTIEMBRE", "ESF"),
        "OCTUBRE_ISR": obtener_horas(df, "OCTUBRE", "ISR")
    }

    print("\n📊 DATOS OBTENIDOS")

    for clave, valor in horas.items():
        print(f"   {clave:<20}: {valor:.2f} horas")

    # -------------------------------------------------
    # 2. CONECTAR CON GOOGLE DRIVE
    # -------------------------------------------------

    print("\n☁️ Conectando con Google Drive...")

    drive = conectar_drive()

    nombre = obtener_nombre_archivo(drive)

    # -------------------------------------------------
    # 3. DESCARGAR XLSX
    # -------------------------------------------------

    print("\n📥 Descargando archivo XLSX...")

    archivo = descargar_excel(drive)

    # -------------------------------------------------
    # 4. MODIFICAR XLSX
    # -------------------------------------------------

    print("\n📝 Modificando celdas...")

    archivo_actualizado = actualizar_excel(archivo, horas)

    print("   ✓ DCG!J14")
    print("   ✓ DCG!K14")
    print("   ✓ DCG!L14")
    print("   ✓ DCG!J15")
    print("   ✓ DCG!K15")
    print("   ✓ DIGDT!J13")
    print("   ✓ DIGDT!K13")

    # -------------------------------------------------
    # 5. SUBIR XLSX ACTUALIZADO
    # -------------------------------------------------

    print("\n📤 Subiendo archivo actualizado a Google Drive...")

    subir_excel(drive, archivo_actualizado)

    print("\n" + "=" * 70)
    print(" ✅ ARCHIVO ACTUALIZADO CORRECTAMENTE")
    print("=" * 70)
    print(f"\n📄 {nombre}")
    print("\nLas demás celdas del archivo no fueron modificadas.")


if __name__ == "__main__":
    actualizar()
